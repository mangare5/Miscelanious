import os
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

# Specify the directory containing your Excel files
root_directory = r"C:\Projects_VS\Miscelanious\TEMU\P_2024_2110_B-H"  # Use raw string for Windows paths


# Walk through all subfolders and files in the root directory
for dirpath, dirnames, filenames in os.walk(root_directory):
    for filename in filenames:
        if filename.endswith(".xlsx") or filename.endswith(".xlsm"):  # Process only supported Excel files
            file_path = os.path.join(dirpath, filename)
            print(f"Processing file: {file_path}")
            try:
                # Load the workbook and select the active worksheet
                workbook = load_workbook(file_path)
                sheet = workbook.active

                # Hide the first row
                sheet.row_dimensions[1].hidden = True

                # Copy content from column D to column E, starting from row 2
                for row in range(2, sheet.max_row + 1):  # Start from row 2
                    cell_value = sheet.cell(row=row, column=4).value  # Column D is 4
                    sheet.cell(row=row, column=5, value=cell_value)  # Column E is 5

                # Save the changes to the Excel file
                workbook.save(file_path)
                print(f"Updated file saved: {file_path}")

            except InvalidFileException:
                print(f"Skipping unsupported or corrupt file: {file_path}")
            except Exception as e:
                print(f"An error occurred while processing {file_path}: {e}")

print("Processing completed for all files.")
